# tests/test_coverage_xlsx.py
"""
Écriture Excel sans dépendance de diag_03_backup_coverage (openpyxl est absent du runtime
Databricks 17.3, constaté sur dev). Le fichier produit est relu avec openpyxl, disponible en local.
"""
from pathlib import Path

import openpyxl

SRC = (Path(__file__).resolve().parent.parent / "notebooks" / "diag_03_backup_coverage.py").read_text(encoding="utf-8")
NS = {}
exec(SRC.split("# >>> XLSX")[1].split("# <<< XLSX")[0], NS)
write_xlsx = NS["write_xlsx"]


def test_xlsx_is_readable_with_sheets_headers_and_values(tmp_path):
    p = tmp_path / "r.xlsx"
    write_xlsx(str(p), {
        "Synthèse": (["Type", "✅ OK", "Total"], [["Table / vue", 3, 4], ["Volume", 0, 1]]),
        "Détail": (["Objet", "Explication"], [["prod.raw.t", "DDL + <données> & \"clone\" ; d'accord"]]),
    })
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Synthèse", "Détail"]
    ws = wb["Synthèse"]
    assert [c.value for c in ws[1]] == ["Type", "✅ OK", "Total"]
    assert [c.value for c in ws[2]] == ["Table / vue", 3, 4]
    assert ws["A1"].font.bold
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:C3"
    assert wb["Détail"]["B2"].value == "DDL + <données> & \"clone\" ; d'accord"


def test_xlsx_handles_none_control_chars_long_names_and_many_columns(tmp_path):
    p = tmp_path / "r.xlsx"
    headers = [f"c{i}" for i in range(30)]
    write_xlsx(str(p), {"Un nom de feuille beaucoup trop long [pour Excel]": (headers, [[None, "a\x01b"] + [1] * 28])})
    wb = openpyxl.load_workbook(p)
    ws = wb.worksheets[0]
    assert len(ws.title) <= 31 and "[" not in ws.title
    assert ws["A2"].value is None
    assert ws["B2"].value == "ab"
    assert ws["AD1"].value == "c29"
