"""Charte commune aux documents Word livrés au client.

Extrait de generate_deployment_guide.py (lui-même adapté de old/generate_guide.py) au moment
où un second générateur a eu besoin des mêmes primitives. Toute modification ici se répercute
sur l'ensemble des documents — c'est l'intention.

Charte : Calibri 11, titres bleu 1F497D, code Courier New 9 sur fond F0F0F0,
tableaux Table Grid à en-tête bleu, notes en italique gris, avertissements en gras rouge.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = (0x1F, 0x49, 0x7D)
GREY = (0x70, 0x70, 0x70)
RED = (0xC0, 0x00, 0x00)


def new_document():
    """Document vierge avec la police de base appliquée."""
    doc = Document()
    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Calibri"
    style_normal.font.size = Pt(11)
    return doc


def _shade(element, fill):
    """Applique un fond de couleur à un paragraphe ou une cellule."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def add_heading(doc, text, level, color=BLUE):
    h = doc.add_heading(text, level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in h.runs:
        if color:
            run.font.color.rgb = RGBColor(*color)
    return h


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(*BLUE)
    _shade(p._p.get_or_add_pPr(), "F0F0F0")
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for run in hdr[i].paragraphs[0].runs:
            run.font.bold = True
        _shade(hdr[i]._tc.get_or_add_tcPr(), "1F497D")
        for para in hdr[i].paragraphs:
            for run in para.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row_data in rows:
        row = table.add_row().cells
        for i, val in enumerate(row_data):
            row[i].text = val
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_note(doc, text, prefix="Note : "):
    p = doc.add_paragraph()
    run = p.add_run(prefix + text)
    run.font.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(*GREY)
    p.paragraph_format.left_indent = Cm(0.5)
    return p


def add_warning(doc, text, prefix="Attention : "):
    p = doc.add_paragraph()
    run = p.add_run(prefix + text)
    run.font.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(*RED)
    p.paragraph_format.left_indent = Cm(0.5)
    return p


def bullet(doc, text):
    return doc.add_paragraph(text, style="List Bullet")


def numbered(doc, text):
    return doc.add_paragraph(text, style="List Number")


def add_title_page(doc, title, subtitle, version, date_str, no_break=False):
    """Page de titre commune : titre bleu, sous-titre, version, mention KeyIT + date."""
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(*BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(version)
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"KeyIT — {date_str}")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    if not no_break:
        doc.add_page_break()
