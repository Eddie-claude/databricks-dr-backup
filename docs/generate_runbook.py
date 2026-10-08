"""Génère DR_Backup_Runbook_Restauration_v4.2.docx depuis runbook_restauration_complete.md.

Le runbook Markdown reste la source unique de la procédure de restauration : ce script le met en
forme avec la charte commune (docs/docx_style.py) pour la livraison. Sous-ensemble Markdown pris en
charge : titres #/##/###/####, tableaux, blocs ```, listes - et 1., citations > (⚠️ = avertissement),
**gras** et `code` en ligne, séparateurs ---.

Usage : python docs/generate_runbook.py
"""

import datetime
import os
import re

from docx.shared import Pt, RGBColor

from docx_style import BLUE, add_code, add_heading, add_note, add_table, add_title_page, add_warning, new_document

SRC = os.path.join(os.path.dirname(os.path.abspath(__file__)), "runbook_restauration_complete.md")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "DR_Backup_Runbook_Restauration_v4.2.docx")

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def plain(text: str) -> str:
    """Texte sans marqueurs Markdown (cellules de tableau, avertissements)."""
    return re.sub(r"\*\*([^*]+)\*\*", r"\1", text).replace("`", "")


def rich(paragraph, text: str) -> None:
    """Ajoute text au paragraphe : **gras** en gras, `code` en Courier New bleu."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).font.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(*BLUE)
        else:
            paragraph.add_run(part)


def table_cells(line: str) -> list:
    return [plain(c.strip()) for c in line.strip().strip("|").split("|")]


def convert(lines: list, doc) -> None:
    i, para = 0, []

    def flush():
        if para:
            rich(doc.add_paragraph(), " ".join(para))
            para.clear()

    while i < len(lines):
        line = lines[i].rstrip("\n")
        s = line.strip()
        if s.startswith("```"):
            flush()
            block, i = [], i + 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip("\n"))
                i += 1
            add_code(doc, "\n".join(block))
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\|[\s:|-]+\|$", lines[i].strip()):
                    rows.append(table_cells(lines[i]))
                i += 1
            add_table(doc, rows[0], rows[1:])
            continue
        elif s.startswith("#"):
            flush()
            level = len(s) - len(s.lstrip("#"))
            if level > 1:                       # le titre # est porté par la page de titre
                add_heading(doc, plain(s.lstrip("#").strip()), min(level - 1, 3))
        elif s.startswith(">"):
            flush()
            text = s.lstrip(">").strip()
            if "⚠️" in text:
                add_warning(doc, plain(text.replace("⚠️", "").strip()), prefix="")
            else:
                add_note(doc, plain(text), prefix="")
        elif re.match(r"^[-*] ", s):
            flush()
            rich(doc.add_paragraph(style="List Bullet"), s[2:])
        elif re.match(r"^\d+\. ", s):
            flush()
            rich(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", s))
        elif s == "---" or not s:
            flush()
        else:
            para.append(s)
        i += 1
    flush()


doc = new_document()
add_title_page(
    doc,
    title="DR Backup Databricks",
    subtitle="Runbook de restauration",
    version="Version 4.2",
    date_str=datetime.date.today().strftime("%d/%m/%Y"),
)
with open(SRC, encoding="utf-8") as f:
    convert(f.readlines(), doc)
doc.save(OUT)
print(f"[OK] Runbook genere : {OUT}")
